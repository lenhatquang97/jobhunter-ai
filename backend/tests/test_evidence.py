import base64
import io
import json

import pytest
from app.documents import extract_document
from app.tools import TOOL_DEFINITIONS, validate_arguments
from conftest import Bridge, Provider, conversation, headers, submission, terminal
from docx import Document


def call(name="get_job", arguments=None, cid="call-1"):
    return {
        "type": "tool.call",
        "id": cid,
        "name": name,
        "arguments": json.dumps(arguments or {"job_id": 11}),
    }


def evidence_provider(text="Build reliable systems. [S1]"):
    return Provider(
        [
            [call()],
            [{"type": "text.delta", "text": "Internal draft"}],
            [{"type": "text.delta", "text": text}],
        ]
    )


async def test_tools_report_exact_citations_and_fresh_history(make_client):
    provider, bridge = evidence_provider(), Bridge()
    client, _ = await make_client(provider, bridge)
    cid = await conversation(client)
    body = submission(context_refs=[{"model": "hr.job", "id": 11}])
    await client.post(f"/v1/conversations/{cid}/messages", json=body)
    result = await terminal(client, body["run_id"])
    assert result["run"]["status"] == "completed", result
    assert [e["agent_id"] for e in result["events"] if e["type"] == "agent.started"] == [
        "Orchestrator",
        "Report",
    ]
    assert len(provider.inputs[-1][1]) == 0
    assert bridge.calls[0][1] == {
        "arguments": {"job_id": 11},
        "context_refs": [{"model": "hr.job", "id": 11}],
    }
    assert any(c[0] == "validate_sources" for c in bridge.calls)
    history = (
        await client.get(f"/v1/conversations/{cid}", headers=headers(capability="fresh"))
    ).json()
    assert history["messages"][-1]["sources"][0]["key"] == "S1"
    assert "Internal draft" not in str(result)
    bridge.state = "unavailable"
    history = (
        await client.get(f"/v1/conversations/{cid}", headers=headers(capability="fresh"))
    ).json()
    assert history["messages"][-1]["visibility"] == "unavailable"
    assert history["messages"][-1]["text"] == ""
    assert history["messages"][-1]["sources"] == []
    replay = (
        await client.get(f"/v1/runs/{body['run_id']}/events", headers=headers(capability="fresh"))
    ).json()
    assert "Build reliable" not in str(replay)
    assert "Synthetic role" not in str(replay)


async def test_no_validation_grant_masks_evidence_and_failed_validation_excludes_prompt(
    make_client,
):
    provider, bridge = evidence_provider(), Bridge()
    client, _ = await make_client(provider, bridge)
    cid = await conversation(client)
    body = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=body)
    await terminal(client, body["run_id"])
    hidden = (await client.get(f"/v1/conversations/{cid}")).json()
    assert hidden["messages"][-1]["text"] == ""
    bridge.fail_validation = True
    provider.turns.append([{"type": "text.delta", "text": "New greeting"}])
    new = submission(text="new")
    await client.post(f"/v1/conversations/{cid}/messages", json=new)
    await terminal(client, new["run_id"])
    assert "Build reliable" not in str(provider.inputs[-1][0])
    assert "secret bridge" not in str((await client.get(f"/v1/conversations/{cid}")).json())


async def test_outdated_history_visible_labeled_but_excluded_from_model(make_client):
    provider, bridge = evidence_provider(), Bridge()
    client, _ = await make_client(provider, bridge)
    cid = await conversation(client)
    first = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=first)
    await terminal(client, first["run_id"])
    bridge.state = "outdated"
    history = (
        await client.get(f"/v1/conversations/{cid}", headers=headers(capability="fresh"))
    ).json()
    assert history["messages"][-1]["visibility"] == "outdated"
    provider.turns.append([{"type": "text.delta", "text": "Hello again"}])
    second = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=second)
    await terminal(client, second["run_id"])
    assert "Build reliable" not in str(provider.inputs[-1][0])


@pytest.mark.parametrize(
    "text", ["Invented [S99]", '"Invented quote" [S1]', "Uncited factual conclusion"]
)
async def test_fabricated_or_missing_citations_fail_closed(make_client, text):
    client, _ = await make_client(evidence_provider(text))
    cid = await conversation(client)
    body = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=body)
    result = await terminal(client, body["run_id"])
    assert result["run"]["status"] == "failed"
    assert result["run"]["error"]["code"] == "invalid_citation"
    assert text not in str(result)


async def test_source_revoked_before_report_delivery_is_not_emitted(make_client):
    bridge = Bridge()

    class Revoker(Provider):
        async def stream(self, messages, tools, role):
            async for event in super().stream(messages, tools, role):
                if role == "Report":
                    bridge.state = "unavailable"
                yield event

    provider = Revoker(evidence_provider().turns)
    client, _ = await make_client(provider, bridge)
    cid = await conversation(client)
    body = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=body)
    final = await terminal(client, body["run_id"])
    assert final["run"]["status"] == "needs_input"
    assert "Build reliable" not in str(final)


async def test_unknown_tool_and_out_of_bounds_arguments_never_callback(make_client):
    for name, args in [
        ("write_applicant", {"id": 1}),
        ("search_jobs", {"query": "", "limit": 21}),
        ("get_job", {"job_id": 11, "sudo": True}),
    ]:
        bridge = Bridge()
        client, _ = await make_client(Provider([[call(name, args)]]), bridge)
        cid = await conversation(client)
        body = submission()
        await client.post(f"/v1/conversations/{cid}/messages", json=body)
        final = await terminal(client, body["run_id"])
        assert final["run"]["status"] == "failed"
        assert not [c for c in bridge.calls if c[0] != "validate_sources"]


async def test_denied_tool_is_failure_not_fabricated_data(make_client):
    bridge = Bridge()
    bridge.result = {
        "ok": False,
        "data": {},
        "sources": [],
        "error": {"code": "access_denied", "message": "Private record 123 exists SECRET"},
    }
    client, _ = await make_client(Provider([[call()]]), bridge)
    cid = await conversation(client)
    body = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=body)
    final = await terminal(client, body["run_id"])
    assert final["run"]["status"] == "failed"
    assert "SECRET" not in str(final)


@pytest.mark.parametrize(
    "name,args",
    [
        ("get_job", {"job_id": True}),
        ("get_job", {"job_id": 0}),
        ("get_applicant", {"applicant_id": "2"}),
        ("search_applicants", {"query": "a", "limit": 2}),
        ("list_documents", {"model": "res.users", "record_id": 2}),
    ],
)
def test_strict_tool_validation(name, args):
    with pytest.raises(ValueError):
        validate_arguments(name, args)


def test_all_tool_schemas_are_strict_and_nullable_fields_required():
    for definition in TOOL_DEFINITIONS:
        assert definition["strict"] is True
        schema = definition["parameters"]
        assert schema["additionalProperties"] is False
        assert set(schema["properties"]) == set(schema["required"])


def document(content, mimetype="text/plain"):
    return {
        "id": 12,
        "name": "synthetic.txt",
        "mimetype": mimetype,
        "content_base64": base64.b64encode(content).decode(),
        "checksum": "abc",
        "write_date": "2026-10-31 12:00:00",
        "res_model": "hr.applicant",
        "res_id": 21,
    }


def test_txt_line_and_docx_paragraph_sources_preserve_exact_locations():
    fragments = extract_document(document(b"first\nsecond\n"))
    assert [(s["locator"], s["excerpt"]) for s in fragments] == [
        ("line:1", "first"),
        ("line:2", "second"),
    ]
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("")
    doc.add_paragraph("Third paragraph")
    data = io.BytesIO()
    doc.save(data)
    fragments = extract_document(
        document(
            data.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )
    assert [(s["locator"], s["excerpt"]) for s in fragments] == [
        ("paragraph:1", "First paragraph"),
        ("paragraph:3", "Third paragraph"),
    ]


@pytest.mark.parametrize(
    "content,mimetype",
    [
        (b"invalid", "application/pdf"),
        (b"invalid", "image/png"),
        (b"\xff", "text/plain"),
        (b"", "text/plain"),
        (b"x" * (10 * 1024 * 1024 + 1), "text/plain"),
    ],
)
def test_unreadable_unsupported_or_oversize_documents_need_input(content, mimetype):
    with pytest.raises(ValueError):
        extract_document(document(content, mimetype))


def test_document_text_is_bounded_and_truncation_explicit():
    fragments = extract_document(document(b"a" * 100001))
    assert sum(len(s["excerpt"]) for s in fragments) <= 100000
    assert fragments[-1]["truncated"] is True


async def test_nonempty_record_without_sources_is_not_accepted_as_untracked_evidence(make_client):
    bridge = Bridge()
    bridge.result = {
        "ok": True,
        "data": {"job": {"id": 11, "name": "Untracked private role"}},
        "sources": [],
    }
    provider = Provider([[call()], [{"type": "text.delta", "text": "Untracked private role"}]])
    client, _ = await make_client(provider, bridge)
    cid = await conversation(client)
    body = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=body)
    result = await terminal(client, body["run_id"])
    assert result["run"]["status"] == "failed"
    assert result["run"]["error"]["code"] == "missing_sources"
    assert "Untracked private role" not in str(result)


async def test_mismatched_attachment_metadata_is_rejected(make_client):
    bridge = Bridge()
    bridge.result = {"ok": True, "data": {"document": document(b"Wrong attachment")}, "sources": []}
    client, _ = await make_client(
        Provider([[call("read_document", {"attachment_id": 99})]]), bridge
    )
    cid = await conversation(client)
    body = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=body)
    result = await terminal(client, body["run_id"])
    assert result["run"]["status"] == "failed"
    assert result["run"]["error"]["code"] == "invalid_tool_result"


async def test_async_document_extraction_isolated_and_exact():
    from app.documents import extract_document_async

    sources = await extract_document_async(document(b"first\nsecond"))
    assert [(s["locator"], s["excerpt"]) for s in sources] == [
        ("line:1", "first"),
        ("line:2", "second"),
    ]


async def test_async_document_extraction_rejects_unreadable_without_parser_details():
    from app.documents import extract_document_async

    with pytest.raises(ValueError, match="readable"):
        await extract_document_async(document(b"PRIVATE_CORRUPT_FILE", "application/pdf"))


def test_pdf_sources_preserve_page_number():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for words in ("First page", "Second page"):
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 100 Td ({words}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    data = io.BytesIO()
    writer.write(data)
    sources = extract_document(document(data.getvalue(), "application/pdf"))
    assert [(s["locator"], s["excerpt"]) for s in sources] == [
        ("page:1", "First page"),
        ("page:2", "Second page"),
    ]


def test_scanned_or_encrypted_pdf_needs_input():
    from pypdf import PdfWriter

    for encrypted in (False, True):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        if encrypted:
            writer.encrypt("synthetic-password")
        data = io.BytesIO()
        writer.write(data)
        with pytest.raises(ValueError):
            extract_document(document(data.getvalue(), "application/pdf"))


def test_many_short_lines_are_bounded_and_explicitly_truncated():
    sources = extract_document(document(b"a\n" * 10000))
    assert len(sources) <= 200
    assert sources[-1]["truncated"] is True


def test_quote_must_match_immediately_attached_source_not_any_later_citation():
    from app.schemas import SafeError
    from app.service import validate_citations

    with pytest.raises(SafeError):
        validate_citations(
            '"First fact" [S2]. Something else [S1]',
            [{"key": "S1", "excerpt": "First fact"}, {"key": "S2", "excerpt": "Second fact"}],
        )


async def test_source_free_answer_cannot_invent_a_citation(make_client):
    client, _ = await make_client(
        Provider([[{"type": "text.delta", "text": "Invented source [S1]"}]])
    )
    cid = await conversation(client)
    body = submission()
    await client.post(f"/v1/conversations/{cid}/messages", json=body)
    result = await terminal(client, body["run_id"])
    assert result["run"]["status"] == "failed"
    assert result["run"]["error"]["code"] == "invalid_citation"
    assert "Invented source" not in str(result)
